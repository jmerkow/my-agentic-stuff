#!/usr/bin/env python3

"""Inspect a PPTX for package integrity, object bounds, and flattened slides."""

from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def inspect_presentation(path: Path, allow_full_slide_images: bool) -> tuple[list[str], list[str]]:
    """Return structural errors and warnings for a presentation."""
    errors: list[str] = []
    warnings: list[str] = []

    if not path.is_file():
        return [f"File does not exist: {path}"], warnings

    if not zipfile.is_zipfile(path):
        return [f"Not a valid ZIP-based Office file: {path}"], warnings

    with zipfile.ZipFile(path) as package:
        corrupt_member = package.testzip()
        if corrupt_member:
            errors.append(f"Corrupt package member: {corrupt_member}")

    try:
        presentation = Presentation(path)
    except Exception as exception:
        return errors + [f"python-pptx could not open the deck: {exception}"], warnings

    if not presentation.slides:
        errors.append("Presentation has no slides")
        return errors, warnings

    slide_width = presentation.slide_width
    slide_height = presentation.slide_height

    for slide_number, slide in enumerate(presentation.slides, start=1):
        if not slide.shapes:
            errors.append(f"Slide {slide_number}: no objects")
            continue

        full_slide_pictures = 0
        for shape in slide.shapes:
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            if shape.left < 0 or shape.top < 0 or right > slide_width or bottom > slide_height:
                errors.append(
                    f"Slide {slide_number}: object {shape.name!r} is outside the slide bounds"
                )

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                width_ratio = shape.width / slide_width
                height_ratio = shape.height / slide_height
                if width_ratio >= 0.95 and height_ratio >= 0.95:
                    full_slide_pictures += 1

        if full_slide_pictures and not allow_full_slide_images:
            errors.append(
                f"Slide {slide_number}: contains a full-slide picture; use native objects or pass "
                "--allow-full-slide-images when flattening is intentional"
            )

        editable_count = sum(
            shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes
        )
        if editable_count == 0:
            warnings.append(f"Slide {slide_number}: has no editable non-picture objects")

        print(
            f"Slide {slide_number}: {len(slide.shapes)} objects, "
            f"{editable_count} non-picture objects"
        )

    return errors, warnings


def main() -> int:
    """Run the command-line inspector."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path, help="Presentation to inspect")
    parser.add_argument(
        "--allow-full-slide-images",
        action="store_true",
        help="Allow slides implemented as one full-slide picture",
    )
    arguments = parser.parse_args()

    errors, warnings = inspect_presentation(arguments.pptx, arguments.allow_full_slide_images)
    for warning in warnings:
        print(f"WARN: {warning}", file=sys.stderr)
    for error in errors:
        print(f"ERROR: {error}", file=sys.stderr)

    if errors:
        return 1
    print(f"OK: {arguments.pptx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())